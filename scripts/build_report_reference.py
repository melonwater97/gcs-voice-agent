"""Build the bundled Pandoc DOCX style reference using python-docx."""
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE


def main():
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(22)
    section.left_margin = section.right_margin = Mm(24)
    for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'Caption']:
        style = document.styles[name]
        style.font.name = 'Arial'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Microsoft YaHei')
        style.font.size = Pt({'Title': 24, 'Heading 1': 16, 'Heading 2': 13, 'Heading 3': 12}.get(name, 10.5))
    document.styles['Normal'].paragraph_format.space_after = Pt(7)
    document.styles['Normal'].paragraph_format.line_spacing = 1.25
    style = document.styles['Table'] if 'Table' in document.styles else document.styles.add_style('Table', WD_STYLE_TYPE.TABLE)
    style.base_style = document.styles['Normal Table']
    style.font.size = Pt(10)
    properties = style.element.find(qn('w:tblPr'))
    if properties is None:
        properties = OxmlElement('w:tblPr'); style.element.append(properties)
    borders = OxmlElement('w:tblBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        edge = OxmlElement('w:' + side)
        for name, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            edge.set(qn('w:' + name), value)
        borders.append(edge)
    properties.append(borders)
    document.add_paragraph('PX4 仿真飞行报告样式参考', 'Title')
    document.add_paragraph('实际报告内容由版本化 Markdown 模板及真实飞行数据生成。')
    target = Path(__file__).resolve().parents[1] / 'desktop/resources/reports/reference.docx'
    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(target)
    print(target)


if __name__ == '__main__':
    main()
